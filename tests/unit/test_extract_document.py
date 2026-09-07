"""Office documents through MarkItDown (task P1-08, spec §6.6, §5.2, §5.3).

Against **real files**, built here rather than checked in: a real .docx zip
written part by part, a spreadsheet from `xlsxwriter`, a deck from
`python-pptx`. A byte string that merely starts with `PK` exercises the error
path and nothing else, and the things worth testing — that a heading survives,
that a spreadsheet becomes a table, that the document's own title reaches
`sources` — only exist inside a file a real converter can read.

Two of these tests are transcriptions of §6.6 rather than of the implementation,
and are the reason the file exists:

* MarkItDown is never asked to fetch anything. §6.6: its `convert()` "is
  intentionally permissive across local files, remote URIs and byte streams",
  so `convert_stream` on already-fetched bytes is the only entry point allowed.
  `test_never_fetches_anything_itself` breaks if someone simplifies that away.
* The converter registry is an allowlist. Dropping it would let a hostile
  archive reach `ZipConverter`, which extracts to a temporary directory — and
  nothing else in the suite would notice.
"""

from __future__ import annotations

import datetime as dt
import io
import logging
import threading
import zipfile
from typing import Any

import pytest
from markitdown import MarkItDown, StreamInfo

from worker.extract import document as doc
from worker.extract.base import TEXT_FLOOR
from worker.extract.chunk import chunk_text
from worker.extract.document import (
    DOCX,
    FAILED_EXTRACTOR,
    PPTX,
    SUPPORTED_MEDIA_TYPES,
    XLSX,
    _w3c_date,
    extract_document,
    supports,
)

try:  # pragma: no cover - import guard
    import xlsxwriter
except ImportError:  # pragma: no cover - exercised only where the lib is absent
    xlsxwriter = None

try:  # pragma: no cover - import guard
    import pptx
except ImportError:  # pragma: no cover - exercised only where the lib is absent
    pptx = None

try:  # pragma: no cover - import guard
    import mammoth
except ImportError:  # pragma: no cover - exercised only where the lib is absent
    mammoth = None

needs_docx = pytest.mark.skipif(mammoth is None, reason="MarkItDown reads .docx through mammoth")
needs_xlsx = pytest.mark.skipif(xlsxwriter is None, reason="needs xlsxwriter to build a real .xlsx")
needs_pptx = pytest.mark.skipif(pptx is None, reason="needs python-pptx to build a real .pptx")

BODY = (
    "Ridership on the Downtown Line rose by eleven per cent over the period, "
    "against a network average of four per cent across the network as a whole. "
    "The evaluation recommends extending the pilot for a further twelve months."
)


# --------------------------------------------------------------------------
# Builders — real files, not fixtures
# --------------------------------------------------------------------------


def build_docx(
    paragraphs: list[tuple[str, str]],
    *,
    title: str | None = None,
    author: str | None = None,
    description: str | None = None,
    language: str | None = None,
    created: str | None = None,
    with_core_properties: bool = True,
) -> bytes:
    """A real .docx: an OOXML zip with the parts Word and mammoth both require.

    Written by hand because `python-docx` is not a dependency of this project
    and adding one to build a test fixture is a poor trade. The parts here are
    the minimum a reader needs — the content-type map, the package
    relationships, the document body, and the metadata part this module reads
    for `title`/`author`/`publication_date`.
    """
    body = "".join(
        f'<w:p><w:pPr><w:pStyle w:val="{style}"/></w:pPr>'
        f'<w:r><w:t xml:space="preserve">{text}</w:t></w:r></w:p>'
        for style, text in paragraphs
    )
    document = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f"<w:body>{body}</w:body></w:document>"
    )
    content_types = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package'
        '.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        '<Override PartName="/word/document.xml" ContentType="application/vnd'
        '.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
        '<Override PartName="/docProps/core.xml" ContentType="application/vnd'
        '.openxmlformats-package.core-properties+xml"/>'
        "</Types>"
    )
    rels = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006'
        '/relationships/officeDocument" Target="word/document.xml"/>'
        '<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/package/2006'
        '/relationships/metadata/core-properties" Target="docProps/core.xml"/>'
        "</Relationships>"
    )
    fields = "".join(
        element
        for element in (
            f"<dc:title>{title}</dc:title>" if title else "",
            f"<dc:creator>{author}</dc:creator>" if author else "",
            f"<dc:description>{description}</dc:description>" if description else "",
            f"<dc:language>{language}</dc:language>" if language else "",
            f'<dcterms:created xsi:type="dcterms:W3CDTF">{created}</dcterms:created>'
            if created
            else "",
        )
    )
    core = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        "<cp:coreProperties"
        ' xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties"'
        ' xmlns:dc="http://purl.org/dc/elements/1.1/"'
        ' xmlns:dcterms="http://purl.org/dc/terms/"'
        ' xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">'
        f"{fields}</cp:coreProperties>"
    )

    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", content_types)
        archive.writestr("_rels/.rels", rels)
        archive.writestr("word/document.xml", document)
        if with_core_properties:
            archive.writestr("docProps/core.xml", core)
    return buffer.getvalue()


def report_docx(**metadata: Any) -> bytes:
    """A document with a heading and enough prose to clear `TEXT_FLOOR`."""
    return build_docx(
        [("Heading1", "Autonomous shuttle pilot"), ("Normal", BODY), ("Normal", BODY)],
        **metadata,
    )


def build_xlsx() -> bytes:
    """A two-sheet workbook, because a sheet name is content MarkItDown emits."""
    buffer = io.BytesIO()
    workbook = xlsxwriter.Workbook(buffer, {"in_memory": True})
    ridership = workbook.add_worksheet("Ridership")
    ridership.write_row(0, 0, ["Corridor", "Boardings"])
    ridership.write_row(1, 0, ["Downtown Line", 121000])
    ridership.write_row(2, 0, ["Circle Line", 84000])
    notes = workbook.add_worksheet("Notes")
    notes.write(0, 0, BODY)
    workbook.close()
    return buffer.getvalue()


def build_pptx() -> bytes:
    """A deck whose slide body carries the sentence the assertions look for."""
    presentation = pptx.Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Autonomous shuttle pilot"
    slide.placeholders[1].text = BODY
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def build_csv() -> bytes:
    return (
        "corridor,boardings\nDowntown Line,121000\nCircle Line,84000\n" + f"# {BODY}\n"
    ).encode()


#: One real sample per supported media type. The completeness probe below reads
#: this rather than a hand-written list, so a media type added to
#: `SUPPORTED_MEDIA_TYPES` without a file that proves it converts fails there.
SAMPLES: dict[str, Any] = {
    DOCX: report_docx,
    XLSX: build_xlsx,
    PPTX: build_pptx,
    "text/csv": build_csv,
    "application/csv": build_csv,
}

#: Which samples can actually be built here. The completeness probe consults it
#: rather than carrying a skip marker per media type.
BUILDABLE = {DOCX: mammoth is not None, XLSX: xlsxwriter is not None, PPTX: pptx is not None}


# --------------------------------------------------------------------------
# Conversion
# --------------------------------------------------------------------------


@needs_docx
async def test_docx_converts_to_markdown_text():
    """The body text survives, and the document is not paginated (§5.3).

    `pages` empty is the load-bearing half: MarkItDown exposes no page
    boundaries for a .docx, so `page_or_offset` on these chunks is a character
    offset. A `pages` tuple invented here would produce citations that open at
    the wrong page and look correct.
    """
    result = await extract_document(report_docx(), media_type=DOCX)

    assert result.extractor == "markitdown"
    assert "eleven per cent" in result.text
    assert result.pages == ()
    assert result.is_paginated is False
    assert result.has_text is True
    assert result.needs_ocr is False


@needs_docx
async def test_docx_structure_becomes_markdown():
    """A heading arrives as a heading, not as an undifferentiated paragraph.

    Structure is what makes chunk boundaries land on section breaks rather than
    mid-sentence (`P2-02`), so losing it is a silent quality loss rather than a
    failure. mammoth's style map is what produces it, and a version bump that
    changed it would show up here.
    """
    result = await extract_document(report_docx(), media_type=DOCX)

    assert "# Autonomous shuttle pilot" in result.text
    assert not result.text.startswith((" ", "\n"))


@needs_docx
async def test_docx_text_is_chunkable():
    """Extraction feeds `chunk_text`, so the output has to be usable by it.

    Cheap end-to-end check that this extractor's contribution to the pipeline is
    the same shape as every other one's: unpaginated text, chunked by offset.
    """
    result = await extract_document(report_docx(), media_type=DOCX)
    chunks = chunk_text(result.text)

    assert chunks
    assert len(result.text) >= TEXT_FLOOR


@needs_xlsx
async def test_xlsx_sheets_become_markdown_tables():
    """A spreadsheet's numbers are the finding; a table is how they stay one.

    Also covers sheet names, which carry meaning ("Ridership") that is lost if
    the converter concatenates cells without them.
    """
    result = await extract_document(build_xlsx(), media_type=XLSX)

    assert result.extractor == "markitdown"
    assert "Ridership" in result.text
    assert "| Downtown Line | 121000 |" in result.text
    assert "Circle Line" in result.text


@needs_pptx
async def test_pptx_slides_become_text():
    """Consultancy findings arrive as decks more often than as reports."""
    result = await extract_document(build_pptx(), media_type=PPTX)

    assert result.extractor == "markitdown"
    assert "Autonomous shuttle pilot" in result.text
    assert "eleven per cent" in result.text


async def test_csv_becomes_a_markdown_table():
    result = await extract_document(build_csv(), media_type="text/csv")

    assert result.extractor == "markitdown"
    assert "| Downtown Line | 121000 |" in result.text


@pytest.mark.parametrize("media_type", sorted(SUPPORTED_MEDIA_TYPES))
async def test_every_supported_media_type_converts_a_real_file(media_type):
    """Completeness probe: every claim in `SUPPORTED_MEDIA_TYPES` is verified.

    The set is the routing table `main.py` reads, so a media type listed there
    that MarkItDown cannot actually convert becomes documents fetched, stored
    and silently left without text. This parametrises over the constant rather
    than over a hardcoded list, so adding a media type without adding a sample
    file fails here instead of in production.
    """
    if not BUILDABLE.get(media_type, True):
        pytest.skip(f"no library available here to build a real {media_type} file")

    assert media_type in SAMPLES, "a supported media type needs a real file proving it converts"
    result = await extract_document(SAMPLES[media_type](), media_type=media_type)

    assert result.extractor == "markitdown"
    assert result.text.strip(), f"{media_type} converted to nothing"


# --------------------------------------------------------------------------
# Metadata — taken from the document, never guessed
# --------------------------------------------------------------------------


@needs_docx
async def test_core_properties_fill_the_citation_fields():
    """A .docx's own metadata reaches `sources` (§5.2).

    MarkItDown returns markdown and nothing else for Office formats, so without
    this every Office source in the corpus would render as a bare URL.
    """
    content = report_docx(
        title="Autonomous shuttle pilot: evaluation",
        author="Land Transport Authority",
        description="Twelve-month evaluation of the shuttle pilot.",
        language="en-SG",
        created="2026-03-14T09:00:00Z",
    )

    result = await extract_document(content, media_type=DOCX)

    assert result.title == "Autonomous shuttle pilot: evaluation"
    assert result.author == "Land Transport Authority"
    assert result.excerpt == "Twelve-month evaluation of the shuttle pilot."
    assert result.language == "en-SG"
    assert result.publication_date == dt.date(2026, 3, 14)


@needs_docx
async def test_a_document_that_states_nothing_gets_no_metadata():
    """Absent stays absent. A guessed title is worse than a missing one.

    The tempting guesses are all here to be refused: the first heading is not
    the title, the filename is not the title, and today is not the publication
    date.
    """
    content = report_docx(with_core_properties=False)

    result = await extract_document(content, media_type=DOCX, filename_hint="lta-report-2026.docx")

    assert result.text
    assert result.title is None
    assert result.author is None
    assert result.publication_date is None
    assert result.language is None
    assert result.excerpt is None


@needs_docx
async def test_blank_metadata_elements_are_not_stored_as_empty_strings():
    """`<dc:title>   </dc:title>` is a document with no title, not a blank one."""
    content = report_docx(title="   ", author=" ")

    result = await extract_document(content, media_type=DOCX)

    assert result.title is None
    assert result.author is None


@needs_docx
async def test_partial_creation_date_is_discarded():
    """OOXML permits `2026` and `2026-03`; a citation needs a real day or none.

    Padding a partial date to the first of the month would put a date the
    document never stated into a corpus whose whole claim is being checkable.
    """
    result = await extract_document(report_docx(created="2026-03"), media_type=DOCX)

    assert result.publication_date is None


@pytest.mark.parametrize(
    ("value", "expected"),
    [
        ("2026-03-14T09:00:00Z", dt.date(2026, 3, 14)),
        ("2026-03-14", dt.date(2026, 3, 14)),
        ("  2026-03-14T09:00:00+08:00  ", dt.date(2026, 3, 14)),
        ("2026-03", None),
        ("2026", None),
        ("2026-13-01", None),  # month 13 parses as digits and is not a date
        ("2026-02-30", None),
        ("14/03/2026", None),
        ("", None),
        (None, None),
    ],
)
def test_w3c_date_takes_only_complete_real_dates(value, expected):
    assert _w3c_date(value) == expected


@needs_docx
async def test_metadata_survives_a_document_with_no_metadata_part():
    """No `docProps/core.xml` at all — the common shape for generated files.

    A `KeyError` here would fail the whole extraction of a perfectly readable
    document, which is the expensive way to be missing a title.
    """
    result = await extract_document(report_docx(with_core_properties=False), media_type=DOCX)

    assert result.extractor == "markitdown"
    assert "eleven per cent" in result.text


# --------------------------------------------------------------------------
# Rejection — a bad document is returned, never raised
# --------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("name", "content"),
    [
        ("empty", b""),
        ("truncated zip", b"PK\x03\x04" + b"\x00" * 40),
        ("not a zip at all", b"This is a plain sentence, served as a Word document."),
        (
            "html mislabelled as docx",
            b"<html><head><title>Error 404</title></head><body><p>Not found.</p></body></html>",
        ),
    ],
)
async def test_unreadable_documents_are_returned_not_raised(name, content):
    """One corrupt .docx must not end a crawl that runs for weeks (§13.4).

    The HTML case is the interesting one. MarkItDown sniffs content and would
    otherwise convert this through its generic HTML converter — putting an error
    page into the corpus as if it were the document. The restricted converter
    registry is what makes it a refusal instead, and this test is what notices
    if that registry is ever widened.
    """
    result = await extract_document(content, media_type=DOCX)

    assert result.extractor == FAILED_EXTRACTOR
    assert result.text == ""
    assert result.has_text is False
    assert result.title is None


async def test_a_zip_that_is_not_a_document_is_refused():
    """A .docx that is really a plain archive does not reach an extractor.

    MarkItDown's default registry would hand this to `ZipConverter`, which
    extracts the members to a temporary directory and converts each one — untrusted
    archive contents written to disk, and a zip bomb's cost paid in disk rather
    than in one refused document.
    """
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("notes.txt", BODY)
        archive.writestr("data.csv", "a,b\n1,2\n")

    result = await extract_document(buffer.getvalue(), media_type=DOCX)

    assert result.extractor == FAILED_EXTRACTOR
    assert result.text == ""


@needs_docx
async def test_a_media_type_this_module_does_not_claim_is_refused():
    """PDF and HTML have their own extractors.

    Claiming them here would hide a routing bug behind a worse conversion —
    MarkItDown would gladly turn a navigation menu into corpus text that
    trafilatura exists to remove.
    """
    for media_type in ("application/pdf", "text/html", "application/zip", ""):
        assert supports(media_type) is False
        result = await extract_document(report_docx(), media_type=media_type)
        assert result.extractor == FAILED_EXTRACTOR


@needs_docx
async def test_an_unsupported_media_type_is_logged_loudly():
    """A routing mistake is a silent corpus gap unless something says so.

    `caplog` is unavailable in this suite (`-p no:logging`), so the assertion
    goes through a handler on the module's own logger.
    """
    records: list[logging.LogRecord] = []
    handler = logging.Handler()
    handler.emit = records.append  # type: ignore[method-assign]
    logger = logging.getLogger("worker.extract.document")
    logger.addHandler(handler)
    try:
        await extract_document(report_docx(), media_type="application/pdf")
    finally:
        logger.removeHandler(handler)

    assert [
        record
        for record in records
        if record.levelno == logging.WARNING and "unsupported media type" in record.getMessage()
    ]


@needs_docx
async def test_a_document_that_converts_to_nothing_is_not_a_failure():
    """Converted-and-empty is not the same state as could-not-be-converted.

    §6.5 makes metadata-only a resting state; a re-extraction sweep (§11.12)
    should retry the broken file and leave the empty one alone, which it cannot
    do if both carry the same extractor name.
    """
    result = await extract_document(build_docx([("Normal", "")]), media_type=DOCX)

    assert result.extractor == "markitdown"
    assert result.text == ""
    assert result.has_text is False


async def test_an_empty_csv_does_not_become_content():
    """A CSV of blank lines converts to an empty markdown table, not to nothing.

    `|  |\\n|  |` is what MarkItDown emits for it — text in the strict sense and
    nothing a graph can be built from. `TEXT_FLOOR` is what keeps it out of the
    corpus, so the assertion is on `has_text` rather than on the string.
    """
    result = await extract_document(b"\n\n", media_type="text/csv")

    assert result.extractor == "markitdown"
    assert result.has_text is False


def test_supports_ignores_case_and_charset_parameters():
    """Servers send `Content-Type: Text/CSV; charset=utf-8`.

    A routing table that misses that form drops those documents with no symptom
    other than a corpus quietly missing them.
    """
    assert supports("Text/CSV; charset=utf-8") is True
    assert supports("  text/csv  ") is True
    assert supports(DOCX.upper()) is True
    assert supports("text/csvx") is False


async def test_a_charset_parameter_survives_routing():
    result = await extract_document(build_csv(), media_type="text/csv; charset=utf-8")

    assert result.extractor == "markitdown"
    assert "Downtown Line" in result.text


# --------------------------------------------------------------------------
# The §6.6 security invariants
# --------------------------------------------------------------------------


@needs_docx
async def test_never_fetches_anything_itself(monkeypatch):
    """`convert_stream` on fetched bytes, never `convert()` on a URL (§6.6).

    AGENTS.md carries this as an invariant because `convert()` and
    `convert_uri()` will read a local file or fetch a remote one with the
    worker's privileges — an SSRF and an arbitrary-file read behind an
    innocuous-looking method. Every fetching entry point is replaced with an
    explosion here; a normal conversion still has to succeed, so the test fails
    the moment someone "simplifies" the call site into one of them.
    """

    def forbidden(*args: Any, **kwargs: Any):
        raise AssertionError("MarkItDown must only ever be given already-fetched bytes")

    for method in ("convert", "convert_local", "convert_uri", "convert_url"):
        monkeypatch.setattr(MarkItDown, method, forbidden)

    result = await extract_document(report_docx(title="Evaluation"), media_type=DOCX)

    assert result.extractor == "markitdown"
    assert "eleven per cent" in result.text
    assert result.title == "Evaluation"


def test_only_the_claimed_converters_are_registered():
    """The registry is an allowlist, and this is the drift test on it.

    MarkItDown's builtins include converters that fetch URLs (YouTube,
    Wikipedia, Bing), shell out to `exiftool`, and extract archives to a
    temporary directory. None of them belongs in a process handling bytes from
    an arbitrary crawled host, and none of the other tests would notice their
    return — a widened registry makes conversions *succeed*, which is exactly
    what makes it easy to miss.
    """
    registered = {
        type(registration.converter).__name__ for registration in doc._markitdown()._converters
    }

    assert registered == {"CsvConverter", "PptxConverter", "XlsxConverter", "DocxConverter"}


def test_the_converter_allowlist_covers_every_supported_media_type():
    """The two lists that can drift apart: what is claimed, and what can run.

    Registering a converter without listing its media type makes a format
    silently unroutable; listing a media type without its converter makes every
    such document fail. Asserted through the converters' own `accepts()`, so it
    stays true if MarkItDown changes which types a converter answers to.
    """
    converters = [registration.converter for registration in doc._markitdown()._converters]
    for media_type in SUPPORTED_MEDIA_TYPES:
        info = StreamInfo(mimetype=media_type, extension=doc._EXTENSIONS[media_type])
        assert any(converter.accepts(io.BytesIO(b""), info) for converter in converters), (
            f"no registered converter accepts {media_type}"
        )


@needs_docx
async def test_conversion_runs_off_the_event_loop(monkeypatch):
    """MarkItDown is synchronous and CPU-bound; the loop runs N fetch lanes.

    A 30-second spreadsheet converted inline stalls every lane, not just its
    own, and the symptom is a crawl that mysteriously slows rather than an
    error. Asserting the conversion happened on another thread is the
    non-flaky way to check that `asyncio.to_thread` is still there.
    """
    threads: list[threading.Thread] = []
    original = doc._convert

    def spy(*args: Any, **kwargs: Any):
        threads.append(threading.current_thread())
        return original(*args, **kwargs)

    monkeypatch.setattr(doc, "_convert", spy)

    result = await extract_document(report_docx(), media_type=DOCX)

    assert result.extractor == "markitdown"
    assert threads and threads[0] is not threading.main_thread()


@needs_docx
async def test_the_filename_hint_cannot_choose_the_converter():
    """The hint comes from a remote server and must not steer routing.

    An attacker who could pick the converter by naming the file `evil.zip`
    would be back at the archive-extraction path the allowlist exists to
    remove. The media type decides; the hint is passed along as context only.
    """
    result = await extract_document(report_docx(), media_type=DOCX, filename_hint="payload.zip")

    assert result.extractor == "markitdown"
    assert "eleven per cent" in result.text


@needs_docx
async def test_xml_metadata_does_not_resolve_external_entities():
    """An XXE in `docProps/core.xml` reads a local file with worker privileges.

    The metadata part is XML from an untrusted archive, and lxml's default
    parser will happily resolve `SYSTEM "file:///etc/passwd"` for it. The
    document must still convert, and the entity must not expand into `title`.
    """
    content = report_docx()
    hostile = (
        '<?xml version="1.0"?>'
        '<!DOCTYPE cp:coreProperties [<!ENTITY xxe SYSTEM "file:///etc/hostname">]>'
        "<cp:coreProperties"
        ' xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties"'
        ' xmlns:dc="http://purl.org/dc/elements/1.1/">'
        "<dc:title>&xxe;</dc:title></cp:coreProperties>"
    )
    rebuilt = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(content)) as source,
        zipfile.ZipFile(rebuilt, "w", zipfile.ZIP_DEFLATED) as target,
    ):
        for name in source.namelist():
            data = hostile.encode() if name == "docProps/core.xml" else source.read(name)
            target.writestr(name, data)

    result = await extract_document(rebuilt.getvalue(), media_type=DOCX)

    assert result.extractor == "markitdown"
    assert "eleven per cent" in result.text
    assert result.title is None or "/" not in result.title
